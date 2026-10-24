"""ورشة PDF المحلية — Local PDF Workshop. Runs on your computer only.

تشغيل:  python pdf_workshop.py
تفتح صفحة في المتصفح على 127.0.0.1 (جهازك فقط). لا يُحفظ أي ملف: كل طلب
يُعالج في مجلد مؤقت يُحذف فوراً بعد إرسال النتيجة، والبرنامج يقفل نفسه
تلقائياً بعد إغلاق الصفحة.

المكتبات: pypdf (دمج/تقسيم/حماية/بيانات وصفية) — PyMuPDF (صور/ضغط/علامة مائية/ترقيم)
— pdf2docx (تحويل Word مع التنسيق) — python-docx (Word للنص العربي).
"""

from __future__ import annotations

import io
import json
import logging
import os
import re
import secrets
import socket
import sys
import tempfile
import threading
import time
import webbrowser
import zipfile
from email.parser import BytesParser
from email.policy import default as email_policy
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, urlparse

try:
    import pymupdf  # PyMuPDF >= 1.24
except ImportError:  # older installs
    import fitz as pymupdf  # type: ignore
from pypdf import PdfReader, PdfWriter
from pypdf.constants import UserAccessPermissions

logging.getLogger("pdf2docx").setLevel(logging.ERROR)
logging.basicConfig(level=logging.ERROR)

MAX_UPLOAD = 600 * 1024 * 1024  # 600 MB per request
IDLE_EXIT_SECONDS = 120          # quit when the page has been closed this long
AR = re.compile(r"[؀-ۿݐ-ݿࢠ-ࣿﭐ-﷿ﹰ-﻿]")
LAT = re.compile(r"[A-Za-z]")


class UserError(Exception):
    """A problem the user can fix (wrong password, bad page range…)."""


_LOCAL = threading.local()


def t(ar: str, en: str) -> str:
    """Pick the message in the language of the current request."""
    return en if getattr(_LOCAL, "lang", "ar") == "en" else ar


# ---------------------------------------------------------------- helpers

def parse_range(spec: str, n: int) -> list[int]:
    """'1-3, 5' -> [0,1,2,4]. Empty -> all pages. Accepts Arabic digits and commas."""
    spec = (spec or "").strip()
    if not spec:
        return list(range(n))
    spec = spec.translate(str.maketrans("٠١٢٣٤٥٦٧٨٩،", "0123456789,"))
    out: list[int] = []
    for part in filter(None, re.split(r"[,\s]+", spec)):
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m:
            raise UserError(t(f"لم أفهم «{part}» في خانة الصفحات. اكتبها مثل: 1-3, 5", f"Could not read \"{part}\" in Pages. Use a format like: 1-3, 5"))
        a, b = int(m[1]), int(m[2] or m[1])
        if a > b:
            a, b = b, a
        for i in range(a, b + 1):
            if not 1 <= i <= n:
                raise UserError(t(f"الصفحة {i} غير موجودة؛ عدد صفحات الملف {n}.", f"Page {i} does not exist; the file has {n} pages."))
            out.append(i - 1)
    return out


def open_reader(path: Path, password: str = "") -> PdfReader:
    try:
        r = PdfReader(str(path))
    except Exception as e:  # noqa: BLE001
        raise UserError(t(f"تعذّر فتح «{path.name}» — قد يكون الملف تالفاً.", f"Could not open \"{path.name}\" — the file may be damaged.")) from e
    if r.is_encrypted:
        if not r.decrypt(password or ""):
            raise UserError(f"«{path.name}» محمي بكلمة مرور. استخدم أداة «فك الحماية» أولاً.")
    return r


def open_fitz(path: Path, password: str = ""):
    try:
        d = pymupdf.open(str(path))
    except Exception as e:  # noqa: BLE001
        raise UserError(t(f"تعذّر فتح «{path.name}».", f"Could not open \"{path.name}\".")) from e
    if d.needs_pass and not d.authenticate(password or ""):
        raise UserError(f"«{path.name}» محمي بكلمة مرور. استخدم أداة «فك الحماية» أولاً.")
    return d


def stem(name: str) -> str:
    s = Path(name).stem or t("مستند", "document")
    return re.sub(r'[\\/:*?"<>|]', "-", s)


def is_rtl(s: str) -> bool:
    return bool(AR.search(s)) and len(AR.findall(s)) >= len(LAT.findall(s))


def human(n: int) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024 or unit == "GB":
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return str(n)


# ---------------------------------------------------------------- operations
# Each op takes (files: list[Path], opts: dict, tmp: Path) and returns
# (filename, bytes, note).
#
# Arabic text is read in the page by pdf.js (vendor/), which keeps word order
# and lam-alef ligatures correct; the page posts the lines here as JSON and
# python-docx builds the Word file. Latin documents go through pdf2docx, which
# also rebuilds tables and layout.

def _docx_from_lines(pages: list[dict], out: Path) -> int:
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = docx.Document()
    st = doc.styles["Normal"]
    st.font.name = "Arial"
    st.element.rPr.rFonts.set(qn("w:cs"), "Arial")
    st.font.size = Pt(12)
    flat = [l for pg in pages for l in pg["lines"]]
    rtl_doc = sum(1 for l in flat if l.get("rtl")) >= len(flat) / 2
    if rtl_doc:
        doc.sections[0]._sectPr.append(OxmlElement("w:bidi"))
    empty = 0
    for k, pg in enumerate(pages):
        if k:
            doc.add_page_break()
        lines = pg["lines"]
        if not lines:
            empty += 1
            lines = [{"text": t(f"[الصفحة {pg['page']}: لا يوجد نص — غالباً صفحة ممسوحة ضوئياً]", f"[Page {pg['page']}: no text found — probably a scanned page]"), "rtl": getattr(_LOCAL, "lang", "ar") != "en", "size": 11}]
        for ln in lines:
            if ln.get("gap"):
                doc.add_paragraph()
            text, rtl = str(ln.get("text", "")), bool(ln.get("rtl"))
            size = Pt(max(8, min(36, round(float(ln.get("size") or 12)))))
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            if rtl:
                p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
            # split Latin chunks out of Arabic lines so Word orders them correctly
            parts = re.split(r"([A-Za-z][\w.,:;/\-@%&+'()]*(?: [A-Za-z][\w.,:;/\-@%&+'()]*)*)", text) if rtl else [text]
            for seg in filter(None, parts):
                run = p.add_run(seg)
                run.font.size = size
                rpr = run._r.get_or_add_rPr()
                szcs = OxmlElement("w:szCs")
                szcs.set(qn("w:val"), str(int(size.pt * 2)))
                rpr.append(szcs)
                if rtl and not LAT.search(seg):
                    rpr.append(OxmlElement("w:rtl"))
    doc.save(str(out))
    return empty


def op_to_docx(files, o, tmp):
    src = files[0]
    out = tmp / "out.docx"
    if o.get("lines"):
        try:
            pages = json.loads(o["lines"])
        except ValueError as e:
            raise UserError(t("بيانات النص غير صالحة.", "Invalid text data.")) from e
        empty = _docx_from_lines(pages, out)
        note = t("نص قابل للتعديل باتجاه صحيح.", "Editable text with correct reading direction.")
        if empty:
            note += t(f" {empty} صفحة بدون نص (ممسوحة ضوئياً).", f" {empty} page(s) have no text (scanned).")
    else:
        pages = parse_range(o.get("range", ""), len(open_reader(src).pages))
        from pdf2docx import Converter
        cv = Converter(str(src))
        try:
            cv.convert(str(out), pages=pages)
        finally:
            cv.close()
        note = t("تم التحويل مع الحفاظ على التنسيق والجداول (pdf2docx).", "Converted with layout and tables kept (pdf2docx).")
    return f"{stem(src.name)}.docx", out.read_bytes(), note


def op_to_images(files, o, tmp):
    src = files[0]
    d = open_fitz(src)
    pages = parse_range(o.get("range", ""), d.page_count)
    dpi = int(o.get("dpi", 150))
    fmt = "jpg" if o.get("fmt") == "jpg" else "png"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for i in pages:
            pix = d[i].get_pixmap(dpi=dpi, alpha=False)
            data = pix.tobytes("jpeg", jpg_quality=90) if fmt == "jpg" else pix.tobytes("png")
            if len(pages) == 1:
                return f"{stem(src.name)}-{i + 1}.{fmt}", data, ""
            z.writestr(f"{stem(src.name)}-{i + 1:03d}.{fmt}", data)
    return f"{stem(src.name)}-{t('صور', 'images')}.zip", buf.getvalue(), t(f"{len(pages)} صورة بدقة {dpi} DPI.", f"{len(pages)} images at {dpi} DPI.")


def op_images_to_pdf(files, o, tmp):
    out = pymupdf.open()
    fit_a4 = o.get("size", "a4") == "a4"
    for f in files:
        try:
            img = pymupdf.open(str(f))
            rect = img[0].rect
            pdf = pymupdf.open("pdf", img.convert_to_pdf())
        except Exception as e:  # noqa: BLE001
            raise UserError(t(f"تعذّر قراءة الصورة «{f.name}».", f"Could not read the image \"{f.name}\".")) from e
        if fit_a4:
            w, h = (842, 595) if rect.width > rect.height else (595, 842)
            page = out.new_page(width=w, height=h)
            page.show_pdf_page(page.rect + (18, 18, -18, -18), pdf, 0)
        else:
            out.insert_pdf(pdf)
    return f"{stem(files[0].name)}.pdf", out.tobytes(garbage=3, deflate=True), t(f"{len(files)} صورة في ملف واحد.", f"{len(files)} images in one file.")


def op_merge(files, o, tmp):
    if len(files) < 2:
        raise UserError(t("أضف ملفين أو أكثر للدمج.", "Add two or more files to merge."))
    w = PdfWriter()
    for f in files:
        w.append(open_reader(f))
    buf = io.BytesIO()
    w.write(buf)
    return f"{stem(files[0].name)}-{t('مدموج', 'merged')}.pdf", buf.getvalue(), t(f"تم دمج {len(files)} ملفات بالترتيب الظاهر.", f"Merged {len(files)} files in the order shown.")


def op_split(files, o, tmp):
    src = files[0]
    r = open_reader(src)
    n = len(r.pages)
    mode = o.get("mode", "range")
    if mode == "range":
        pages = parse_range(o.get("range", ""), n)
        w = PdfWriter()
        for i in pages:
            w.add_page(r.pages[i])
        buf = io.BytesIO()
        w.write(buf)
        return f"{stem(src.name)}-{t('مستخرج', 'extract')}.pdf", buf.getvalue(), t(f"{len(pages)} صفحة.", f"{len(pages)} pages.")
    every = 1 if mode == "each" else max(1, int(o.get("every", 1) or 1))
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as z:
        for start in range(0, n, every):
            w = PdfWriter()
            for i in range(start, min(n, start + every)):
                w.add_page(r.pages[i])
            b = io.BytesIO()
            w.write(b)
            end = min(n, start + every)
            label = f"{start + 1}" if end - start == 1 else f"{start + 1}-{end}"
            z.writestr(f"{stem(src.name)}-{label}.pdf", b.getvalue())
    return f"{stem(src.name)}-{t('مقسم', 'split')}.zip", zbuf.getvalue(), ""


def op_pages(files, o, tmp):
    src = files[0]
    r = open_reader(src)
    n = len(r.pages)
    order = parse_range(o.get("order", ""), n)
    delete = set(parse_range(o["delete"], n)) if o.get("delete", "").strip() else set()
    rotate = set(parse_range(o["rotate"], n)) if o.get("rotate", "").strip() else set()
    angle = int(o.get("angle", 90))
    w = PdfWriter()
    for i in order:
        if i in delete:
            continue
        page = w.add_page(r.pages[i])
        if i in rotate:
            page.rotate(angle)
    if not len(w.pages):
        raise UserError(t("حذفت كل الصفحات — لازم تبقى صفحة واحدة على الأقل.", "That deletes every page — at least one page must remain."))
    buf = io.BytesIO()
    w.write(buf)
    return f"{stem(src.name)}-{t('معدل', 'edited')}.pdf", buf.getvalue(), t(f"الناتج {len(w.pages)} صفحة.", f"Result: {len(w.pages)} pages.")


def op_protect(files, o, tmp):
    src = files[0]
    pw = o.get("password", "")
    if len(pw) < 4:
        raise UserError(t("اكتب كلمة مرور من 4 خانات على الأقل.", "Enter a password of at least 4 characters."))
    perms = UserAccessPermissions(0)
    if o.get("allow_print") == "1":
        perms |= UserAccessPermissions.PRINT | UserAccessPermissions.PRINT_TO_REPRESENTATION
    if o.get("allow_copy") == "1":
        perms |= UserAccessPermissions.EXTRACT | UserAccessPermissions.EXTRACT_TEXT_AND_GRAPHICS
    w = PdfWriter(clone_from=open_reader(src))
    w.encrypt(user_password=pw, owner_password=o.get("owner") or secrets.token_urlsafe(24),
              algorithm="AES-256", permissions_flag=perms)
    buf = io.BytesIO()
    w.write(buf)
    return f"{stem(src.name)}-{t('محمي', 'protected')}.pdf", buf.getvalue(), t("تشفير AES-256. احفظ كلمة المرور؛ لا يمكن استرجاعها.", "AES-256 encryption. Keep the password safe; it cannot be recovered.")


def op_unlock(files, o, tmp):
    src = files[0]
    r = PdfReader(str(src))
    if not r.is_encrypted:
        raise UserError(t("الملف غير محمي أصلاً.", "This file is not password-protected."))
    if not r.decrypt(o.get("password", "")):
        raise UserError(t("كلمة المرور غير صحيحة.", "Wrong password."))
    w = PdfWriter(clone_from=r)
    buf = io.BytesIO()
    w.write(buf)
    return f"{stem(src.name)}-{t('مفتوح', 'unlocked')}.pdf", buf.getvalue(), t("أزيلت الحماية.", "Protection removed.")


def op_compress(files, o, tmp):
    src = files[0]
    before = src.stat().st_size
    d = open_fitz(src)
    level = o.get("level", "medium")
    if level != "light" and hasattr(d, "rewrite_images"):
        dpi, q = (150, 75) if level == "medium" else (100, 55)
        d.rewrite_images(dpi_threshold=dpi + 20, dpi_target=dpi, quality=q, lossy=True, lossless=True, bitonal=True, color=True, gray=True)
    data = d.tobytes(garbage=4, deflate=True, deflate_images=True, deflate_fonts=True, clean=True, use_objstms=1)
    after = len(data)
    if after >= before:
        return f"{stem(src.name)}-{t('مضغوط', 'compressed')}.pdf", src.read_bytes(), t(f"الملف مضغوط أصلاً ({human(before)})؛ أعدته كما هو.", f"Already compact ({human(before)}); returned unchanged.")
    return f"{stem(src.name)}-{t('مضغوط', 'compressed')}.pdf", data, t(f"من {human(before)} إلى {human(after)} (−{100 - after * 100 // before}%).", f"{human(before)} → {human(after)} (−{100 - after * 100 // before}%).")


def _html_text(text: str, size: float, color: str, align: str = "center") -> str:
    from html import escape
    d = "rtl" if is_rtl(text) else "ltr"
    return (f'<div dir="{d}" style="font-family:sans-serif;font-size:{size}px;color:{color};'
            f'text-align:{align};font-weight:bold">{escape(text)}</div>')


def op_watermark(files, o, tmp):
    src = files[0]
    d = open_fitz(src)
    text = (o.get("text") or "").strip() or t("سري", "CONFIDENTIAL")
    opacity = max(0.05, min(1.0, float(o.get("opacity", 20)) / 100))
    angle = int(o.get("angle", 45))
    wm = pymupdf.open()
    # wide enough that the text stays on one line; show_pdf_page scales it to the page
    wp = wm.new_page(width=max(700, len(text) * (38 if is_rtl(text) else 52)), height=140)
    wp.insert_htmlbox(wp.rect, _html_text(text, 72, "#777777"), opacity=opacity)
    pages = parse_range(o.get("range", ""), d.page_count)
    for i in pages:
        p = d[i]
        r = p.rect
        side = min(r.width, r.height) * 0.9
        box = pymupdf.Rect(r.x0 + (r.width - side) / 2, r.y0 + (r.height - side) / 2,
                           r.x0 + (r.width + side) / 2, r.y0 + (r.height + side) / 2)
        p.show_pdf_page(box, wm, 0, rotate=angle - p.rotation, overlay=True)
    return f"{stem(src.name)}-{t('علامة', 'watermarked')}.pdf", d.tobytes(garbage=3, deflate=True), ""


def op_numbers(files, o, tmp):
    src = files[0]
    d = open_fitz(src)
    fmt = o.get("fmt", "ar")
    start = int(o.get("start", 1) or 1)
    total = d.page_count + start - 1
    for i, p in enumerate(d):
        n = i + start
        label = {"ar": f"صفحة {n} من {total}", "en": f"Page {n} of {total}", "num": f"{n}"}.get(fmt, str(n))
        r = p.rect
        box = pymupdf.Rect(r.x0 + 36, r.y1 - 34, r.x1 - 36, r.y1 - 12)
        if p.rotation:
            box = box * p.derotation_matrix
        p.insert_htmlbox(box, _html_text(label, 10, "#444444").replace("font-weight:bold", "font-weight:normal"),
                         rotate=p.rotation)
    return f"{stem(src.name)}-{t('مرقم', 'numbered')}.pdf", d.tobytes(garbage=3, deflate=True), ""


def op_metadata(files, o, tmp):
    src = files[0]
    r = open_reader(src)
    found = {k.lstrip("/"): str(v) for k, v in (r.metadata or {}).items() if str(v).strip()}
    w = PdfWriter(clone_from=r)
    w.metadata = None
    if "/Metadata" in w._root_object:
        del w._root_object["/Metadata"]
    buf = io.BytesIO()
    w.write(buf)
    listed = t("، ", ", ").join(f"{k}={v[:40]}" for k, v in found.items())
    note = t("أزيلت البيانات الوصفية: ", "Metadata removed: ") + (listed if found else t("لم يكن فيه بيانات تُذكر.", "there was none to speak of."))
    return f"{stem(src.name)}-{t('نظيف', 'clean')}.pdf", buf.getvalue(), note


OPS = {
    "to_docx": op_to_docx, "to_images": op_to_images,
    "images_to_pdf": op_images_to_pdf, "merge": op_merge, "split": op_split,
    "pages": op_pages, "protect": op_protect, "unlock": op_unlock,
    "compress": op_compress, "watermark": op_watermark, "numbers": op_numbers,
    "metadata": op_metadata,
}


# ---------------------------------------------------------------- server

TOKEN = secrets.token_urlsafe(24)
_HERE = Path(__file__).resolve().parent
VENDOR = next((d for d in (_HERE / "vendor", _HERE.parent / "vendor") if d.is_dir()), _HERE / "vendor")
STATE = {"last_ping": time.time(), "port": 0}


class Handler(BaseHTTPRequestHandler):
    server_version = "PDFWorkshop"
    sys_version = ""

    def log_message(self, *args):  # never log file names or paths
        pass

    def _host_ok(self) -> bool:
        # blocks DNS-rebinding: only our exact loopback host is accepted
        return self.headers.get("Host") in (f"127.0.0.1:{STATE['port']}", f"localhost:{STATE['port']}")

    def _send(self, code, body: bytes, ctype="application/json; charset=utf-8", extra=None):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Referrer-Policy", "no-referrer")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _json(self, code, obj):
        self._send(code, json.dumps(obj, ensure_ascii=False).encode())

    def do_GET(self):
        if not self._host_ok():
            return self._send(403, b"forbidden", "text/plain")
        u = urlparse(self.path)
        if u.path == "/" and parse_qs(u.query).get("t", [""])[0] == TOKEN:
            page = PAGE.replace("__TOKEN__", TOKEN).encode()
            return self._send(200, page, "text/html; charset=utf-8", {
                "Content-Security-Policy": "default-src 'none'; script-src 'self' 'unsafe-inline'; worker-src 'self' blob:; "
                                           "style-src 'unsafe-inline'; font-src data:; connect-src 'self' data: blob:; "
                                           "img-src 'self' data: blob:; frame-ancestors 'none'; form-action 'none'",
                "X-Frame-Options": "DENY"})
        if u.path in ("/vendor/pdf.min.js", "/vendor/pdf.worker.min.js"):
            return self._send(200, (VENDOR / u.path.rsplit("/", 1)[1]).read_bytes(), "text/javascript; charset=utf-8")
        return self._send(404, b"not found", "text/plain")

    def do_POST(self):
        if not self._host_ok() or self.headers.get("X-Token") != TOKEN:
            return self._send(403, b"forbidden", "text/plain")
        path = urlparse(self.path).path
        if path == "/api/ping":
            STATE["last_ping"] = time.time()
            return self._json(200, {"ok": True})
        if path == "/api/quit":
            self._json(200, {"ok": True})
            threading.Thread(target=lambda: (time.sleep(0.3), os._exit(0)), daemon=True).start()
            return
        if path != "/api/run":
            return self._send(404, b"not found", "text/plain")
        length = int(self.headers.get("Content-Length") or 0)
        _LOCAL.lang = "en" if self.headers.get("X-Lang") == "en" else "ar"
        if length > MAX_UPLOAD:
            return self._json(413, {"error": t("حجم الملفات أكبر من 600 ميغابايت.", "The files are larger than 600 MB.")})
        raw = self.rfile.read(length)
        STATE["last_ping"] = time.time()
        with tempfile.TemporaryDirectory(prefix="pdfws-") as td:
            tmp = Path(td)
            try:
                op, opts, files = self._parse(raw, tmp)
                if op not in OPS:
                    raise UserError(t("أداة غير معروفة.", "Unknown tool."))
                if not files:
                    raise UserError(t("أضف ملفاً أولاً.", "Add a file first."))
                name, data, note = OPS[op](files, opts, tmp)
            except UserError as e:
                return self._json(400, {"error": str(e)})
            except Exception as e:  # noqa: BLE001
                return self._json(500, {"error": t("حدث خطأ أثناء المعالجة: ", "Processing failed: ") + f"{type(e).__name__}: {e}"})
            finally:
                raw = None
        self._send(200, data, "application/octet-stream", {
            "X-Filename": quote(name), "X-Note": quote(note or "")})

    def _parse(self, raw: bytes, tmp: Path):
        # rebuild a MIME message from the multipart body
        ctype = self.headers.get("Content-Type", "")
        if not ctype.startswith("multipart/form-data"):
            raise UserError(t("طلب غير صالح.", "Invalid request."))
        msg_head = f"Content-Type: {ctype}\r\n\r\n".encode()
        msg = BytesParser(policy=email_policy).parsebytes(msg_head + raw)
        opts: dict[str, str] = {}
        files: list[Path] = []
        for k, part in enumerate(msg.iter_parts()):
            name = part.get_param("name", header="content-disposition")
            fname = part.get_filename()
            payload = part.get_payload(decode=True) or b""
            if fname:
                safe = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", Path(fname).name)[-120:] or "file"
                p = tmp / f"{k:03d}" / safe
                p.parent.mkdir()
                p.write_bytes(payload)
                files.append(p)
            elif name:
                opts[name] = payload.decode("utf-8", "replace")
        return opts.pop("op", ""), opts, files



def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def watchdog():
    while True:
        time.sleep(10)
        if time.time() - STATE["last_ping"] > IDLE_EXIT_SECONDS:
            print("أُغلقت الصفحة — تم إيقاف البرنامج. / Page closed — stopped.")
            os._exit(0)


def main():
    port = int(os.environ.get("PDFWS_PORT") or free_port())
    STATE["port"] = port
    httpd = ThreadingHTTPServer(("127.0.0.1", port), Handler)
    url = f"http://127.0.0.1:{port}/?t={TOKEN}"
    print("ورشة PDF تعمل على جهازك فقط. / PDF Workshop is running on this computer only.")
    print("افتح الرابط إذا ما انفتح تلقائياً / Open this link if it did not open:\n  " + url)
    print("للإيقاف / To stop: close this window or press Ctrl+C")
    if os.environ.get("PDFWS_NO_BROWSER") != "1":
        threading.Timer(0.6, lambda: webbrowser.open(url)).start()
        threading.Thread(target=watchdog, daemon=True).start()
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        pass


# ---------------------------------------------------------------- page

PAGE = r"""<!doctype html>
<html lang="ar" dir="rtl"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Local PDF Workshop</title>
<style>
:root{--bg:#eef1f0;--surface:#fff;--sunken:#e3e8e6;--ink:#16211e;--muted:#5a6a65;--line:#cfd8d5;--accent:#0d6b58;--accent-ink:#fff;--accent-soft:#d6ebe5;--danger:#b3261e;--ok:#0d6b58;
--f:"IBM Plex Sans Arabic","Segoe UI","Tahoma",system-ui,sans-serif}
@media (prefers-color-scheme:dark){:root{--bg:#0f1513;--surface:#18201d;--sunken:#111816;--ink:#e6eeeb;--muted:#97a8a2;--line:#2b3733;--accent:#43b89c;--accent-ink:#06231c;--accent-soft:#1b3a32;--danger:#f2867d;--ok:#43b89c;color-scheme:dark}}
*{box-sizing:border-box}[hidden]{display:none!important}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.6 var(--f)}
.wrap{max-width:1100px;margin:0 auto;padding:20px 16px 48px;display:flex;flex-direction:column;gap:16px}
header{display:flex;flex-wrap:wrap;justify-content:space-between;align-items:center;gap:12px}
h1{margin:0;font-size:1.5rem}.sub{margin:0;color:var(--muted);font-size:.9rem}
.row{display:flex;gap:8px;align-items:center;flex-wrap:wrap}
.shield{background:var(--accent-soft);border-radius:999px;padding:6px 14px;font-size:.85rem;font-weight:600}
.bench{display:grid;grid-template-columns:260px minmax(0,1fr);gap:16px;align-items:start}
@media(max-width:760px){.bench{grid-template-columns:1fr}}
.panel{background:var(--surface);border:1px solid var(--line);border-radius:14px;min-width:0}
nav{padding:8px;display:flex;flex-direction:column;gap:2px}
nav h4{margin:10px 10px 4px;font-size:.75rem;color:var(--muted);letter-spacing:.04em}
nav button{all:unset;cursor:pointer;padding:8px 12px;border-radius:8px;font-size:.92rem}
nav button:hover{background:var(--sunken)}nav button[aria-current="true"]{background:var(--accent-soft);font-weight:600}
nav button:focus-visible{outline:2px solid var(--accent)}
main.panel{padding:20px;display:flex;flex-direction:column;gap:16px}
h2{margin:0;font-size:1.2rem}.desc{margin:0;color:var(--muted)}
.drop{border:2px dashed var(--line);border-radius:12px;padding:18px;display:flex;flex-wrap:wrap;gap:12px;align-items:center;justify-content:space-between}
.drop.over{border-color:var(--accent);background:var(--accent-soft)}
.btn,button.btn{font:inherit;cursor:pointer;border:1px solid var(--line);background:var(--surface);color:var(--ink);border-radius:8px;padding:8px 14px}
.btn.primary{background:var(--accent);border-color:var(--accent);color:var(--accent-ink);font-weight:600}
.btn:disabled{opacity:.5;cursor:wait}
.files{display:flex;flex-direction:column;gap:6px;margin:0;padding:0;list-style:none}
.files li{display:flex;align-items:center;gap:8px;background:var(--sunken);border-radius:8px;padding:6px 10px;font-size:.9rem}
.files li span{flex:1;min-width:0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.files li small{color:var(--muted);font-variant-numeric:tabular-nums}
.files li button{all:unset;cursor:pointer;padding:0 6px;color:var(--muted)}
.opts{display:grid;grid-template-columns:repeat(auto-fill,minmax(220px,1fr));gap:12px}
.opts label{display:flex;flex-direction:column;gap:4px;font-size:.85rem;color:var(--muted)}
.opts label.chk{flex-direction:row;align-items:center;color:var(--ink)}
input,select{font:inherit;color:var(--ink);background:var(--surface);border:1px solid var(--line);border-radius:8px;padding:7px 10px}
.run{display:flex;flex-wrap:wrap;gap:12px;align-items:center}
.msg{font-size:.9rem}.msg.err{color:var(--danger)}.msg.ok{color:var(--ok)}
footer{color:var(--muted);font-size:.82rem;display:flex;justify-content:space-between;flex-wrap:wrap;gap:8px}
</style></head><body>
<div class="wrap">
<header><div><h1 id="h1"></h1><p class="sub" id="sub"></p></div>
<div class="row"><div class="shield" id="shield"></div><button class="btn" id="lang" type="button"></button></div></header>
<div class="bench">
<nav class="panel" id="nav"></nav>
<main class="panel">
 <div><h2 id="title"></h2><p class="desc" id="desc"></p></div>
 <div class="drop" id="drop"><div><strong id="dropTitle"></strong><div class="desc" id="dropHint"></div></div>
  <label class="btn primary" for="pick" id="pickLbl"></label><input type="file" id="pick" hidden></div>
 <ul class="files" id="files"></ul>
 <div class="opts" id="opts"></div>
 <div class="run"><button class="btn primary" id="go" type="button"></button><span class="msg" id="msg" role="status"></span></div>
</main></div>
<footer><span id="foot"></span><button class="btn" id="quit" type="button"></button></footer>
</div>
<script src="/vendor/pdf.min.js"></script>
<script>
const TOKEN="__TOKEN__";
pdfjsLib.GlobalWorkerOptions.workerSrc="/vendor/pdf.worker.min.js";
let L=/^ar/i.test(navigator.language||"")?"ar":"en";
const T=(ar,en)=>L==="en"?en:ar;
const UI={
 h1:["ورشة PDF المحلية","Local PDF Workshop"],sub:["تعمل على جهازك فقط · pypdf + PyMuPDF + pdf2docx","Runs on this computer only · pypdf + PyMuPDF + pdf2docx"],
 shield:["🔒 127.0.0.1 — لا رفع ولا حفظ","🔒 127.0.0.1 — nothing uploaded or stored"],lang:["English","عربي"],
 pickLbl:["اختيار ملفات","Choose files"],go:["تنفيذ وحفظ","Run and save"],quit:["إيقاف البرنامج","Stop the program"],
 foot:["الملفات تُعالج في مجلد مؤقت ويُحذف فوراً. البرنامج يقفل نفسه بعد إغلاق الصفحة بدقيقتين.","Files are processed in a temporary folder that is deleted right away. The program stops two minutes after you close this page."]};
const ph=(ar,en)=>T(ar,en);
const rangeF=()=>`<label>${T("الصفحات (فارغ = الكل)","Pages (empty = all)")}<input name="range" placeholder="${ph("مثلاً 1-3, 5","e.g. 1-3, 5")}" dir="ltr"></label>`;
const sel=(name,label,opts,def)=>`<label>${label}<select name="${name}">${opts.map(([v,t])=>`<option value="${v}"${v===def?" selected":""}>${t}</option>`).join("")}</select></label>`;
const G={conv:["تحويل","Convert"],org:["تنظيم","Organize"],sec:["حماية وتعديل","Protect & edit"]};
const TOOLS=[
 {g:"conv",id:"to_docx",t:["PDF إلى Word","PDF to Word"],d:["العربي يتحول لنص قابل للتعديل باتجاه صحيح، والإنجليزي يحافظ على التنسيق والجداول.","English keeps its layout and tables; Arabic becomes clean editable right-to-left text."],
  o:()=>rangeF()+sel("engine",T("طريقة التحويل","Method"),[["auto",T("تلقائي (حسب اللغة)","Automatic (by language)")],["layout",T("الحفاظ على التنسيق والجداول (pdf2docx)","Keep layout and tables (pdf2docx)")],["text",T("نص نظيف قابل للتعديل","Clean editable text")]],"auto")},
 {g:"conv",id:"to_text",t:["PDF إلى نص","PDF to text"],d:["استخراج النص لملف TXT بترتيب صحيح للعربي والإنجليزي.","Extract the text to a TXT file, in the right order for Arabic and English."],o:rangeF},
 {g:"conv",id:"to_images",t:["PDF إلى صور","PDF to images"],d:["صورة لكل صفحة، وتنزل في ملف ZIP.","One image per page, saved in a ZIP file."],
  o:()=>rangeF()+sel("dpi",T("الدقة (DPI)","Resolution (DPI)"),[["100","100"],["150","150"],["200","200"],["300","300"]],"150")+sel("fmt",T("الصيغة","Format"),[["png","PNG"],["jpg","JPG"]],"png")},
 {g:"conv",id:"images_to_pdf",t:["صور إلى PDF","Images to PDF"],multi:1,img:1,d:["اجمع الصور في ملف PDF واحد بالترتيب.","Combine images into one PDF, in order."],
  o:()=>sel("size",T("حجم الصفحة","Page size"),[["a4","A4"],["orig",T("حجم الصورة","Image size")]],"a4")},
 {g:"org",id:"merge",t:["دمج ملفات","Merge files"],multi:1,d:["رتّب الملفات بالأسهم ثم ادمجها في ملف واحد.","Order the files with the arrows, then merge them into one."],o:()=>""},
 {g:"org",id:"split",t:["تقسيم واستخراج","Split and extract"],d:["استخرج صفحات محددة، أو قسّم الملف.","Extract chosen pages, or split the file."],
  o:()=>sel("mode",T("الطريقة","Method"),[["range",T("استخراج صفحات في ملف واحد","Extract pages into one file")],["each",T("كل صفحة في ملف","One file per page")],["every",T("كل عدد صفحات في ملف","Every N pages per file")]],"range")+rangeF()+`<label>${T("عدد الصفحات لكل ملف","Pages per file")}<input name="every" type="number" min="1" value="2"></label>`},
 {g:"org",id:"pages",t:["ترتيب وحذف وتدوير","Reorder, delete, rotate"],d:["اكتب ترتيباً جديداً، أو صفحات للحذف أو التدوير.","Type a new order, or pages to delete or rotate."],
  o:()=>`<label>${T("الترتيب الجديد (فارغ = كما هو)","New order (empty = unchanged)")}<input name="order" placeholder="${ph("مثلاً 3,1,2,4-10","e.g. 3,1,2,4-10")}" dir="ltr"></label><label>${T("حذف الصفحات","Delete pages")}<input name="delete" placeholder="${ph("مثلاً 2, 7-8","e.g. 2, 7-8")}" dir="ltr"></label><label>${T("تدوير الصفحات","Rotate pages")}<input name="rotate" placeholder="${ph("مثلاً 1-3","e.g. 1-3")}" dir="ltr"></label>`+sel("angle",T("زاوية التدوير","Rotation"),[["90",T("90° يمين","90° clockwise")],["270",T("90° يسار","90° counter-clockwise")],["180","180°"]],"90")},
 {g:"sec",id:"protect",t:["حماية بكلمة مرور","Password protect"],d:["تشفير AES-256 مع التحكم بالطباعة والنسخ.","AES-256 encryption with print and copy controls."],
  o:()=>`<label>${T("كلمة المرور","Password")}<input name="password" type="password" autocomplete="new-password"></label><label class="chk"><input type="checkbox" name="allow_print" value="1" checked> ${T("السماح بالطباعة","Allow printing")}</label><label class="chk"><input type="checkbox" name="allow_copy" value="1"> ${T("السماح بنسخ النص","Allow copying text")}</label>`},
 {g:"sec",id:"unlock",t:["فك الحماية","Remove password"],d:["إزالة كلمة المرور من ملف تعرف كلمته.","Remove the password from a file you know the password to."],
  o:()=>`<label>${T("كلمة المرور الحالية","Current password")}<input name="password" type="password" autocomplete="off"></label>`},
 {g:"sec",id:"metadata",t:["إزالة البيانات الوصفية","Remove metadata"],d:["يحذف اسم الكاتب والبرنامج والتواريخ قبل إرسال الملف للخارج.","Removes author, software and dates before you send a file outside."],o:()=>""},
 {g:"sec",id:"watermark",t:["علامة مائية","Watermark"],d:["نص قطري على الصفحات، يدعم العربي والإنجليزي.","Diagonal text across the pages, in Arabic or English."],
  o:()=>`<label>${T("النص","Text")}<input name="text" value="${T("سري — للاستخدام الداخلي","CONFIDENTIAL — INTERNAL USE")}"></label><label>${T("الشفافية %","Opacity %")}<input name="opacity" type="number" min="5" max="100" value="20"></label><label>${T("الزاوية","Angle")}<input name="angle" type="number" value="45"></label>`+rangeF()},
 {g:"sec",id:"numbers",t:["ترقيم الصفحات","Page numbers"],d:["رقم أسفل كل صفحة.","A number at the bottom of every page."],
  o:()=>sel("fmt",T("الشكل","Format"),[["ar","صفحة 1 من 10"],["en","Page 1 of 10"],["num","1"]],L)+`<label>${T("يبدأ من","Start at")}<input name="start" type="number" min="1" value="1"></label>`},
 {g:"sec",id:"compress",t:["ضغط الحجم","Compress"],d:["تقليل حجم الملف لإرساله بالإيميل.","Make the file smaller for email."],
  o:()=>sel("level",T("المستوى","Level"),[["light",T("خفيف (بدون مساس بالصور)","Light (images untouched)")],["medium",T("متوسط (صور 150 DPI)","Medium (images at 150 DPI)")],["strong",T("قوي (صور 100 DPI)","Strong (images at 100 DPI)")]],"medium")}
];
const $=s=>document.querySelector(s);let cur=null,files=[];
const tx=a=>a[L==="en"?1:0];
function applyLang(){document.documentElement.lang=L;document.documentElement.dir=L==="ar"?"rtl":"ltr";
 document.title=tx(UI.h1);for(const k in UI){const e=document.getElementById(k);if(e)e.textContent=tx(UI[k])}nav();pick(cur?cur.id:"to_docx",true)}
$("#lang").onclick=()=>{L=L==="ar"?"en":"ar";applyLang()};
function nav(){let h="",g="";for(const t of TOOLS){if(t.g!==g){g=t.g;h+=`<h4>${tx(G[g])}</h4>`}h+=`<button type="button" data-id="${t.id}">${tx(t.t)}</button>`}$("#nav").innerHTML=h;
 document.querySelectorAll("#nav button").forEach(b=>b.onclick=()=>pick(b.dataset.id))}
function pick(id,keep){cur=TOOLS.find(t=>t.id===id);if(!keep)files=[];document.querySelectorAll("#nav button").forEach(b=>b.setAttribute("aria-current",b.dataset.id===id));
 $("#title").textContent=tx(cur.t);$("#desc").textContent=tx(cur.d);$("#opts").innerHTML=cur.o();$("#msg").textContent="";
 const p=$("#pick");p.multiple=!!cur.multi;p.accept=cur.img?"image/*":"application/pdf,.pdf";p.value="";
 $("#dropTitle").textContent=cur.img?T("اسحب الصور هنا","Drop images here"):cur.multi?T("اسحب الملفات هنا","Drop PDF files here"):T("اسحب ملف PDF هنا","Drop a PDF here");
 $("#dropHint").textContent=cur.multi?T("تقدر تضيف أكثر من ملف وترتبهم.","You can add several files and reorder them."):T("ملف واحد.","One file.");renderFiles()}
function size(n){return n>1048576?(n/1048576).toFixed(1)+" MB":Math.ceil(n/1024)+" KB"}
function renderFiles(){$("#files").innerHTML="";files.forEach((f,i)=>{const li=document.createElement("li");
 const up=T("تقديم","Move up"),dn=T("تأخير","Move down"),rm=T("إزالة","Remove");
 li.innerHTML=`<span></span><small>${size(f.size)}</small>`+(cur.multi?`<button title="${up}" aria-label="${up}">▲</button><button title="${dn}" aria-label="${dn}">▼</button>`:"")+`<button title="${rm}" aria-label="${rm}">✕</button>`;
 li.querySelector("span").textContent=f.name;const bs=li.querySelectorAll("button");
 if(cur.multi){bs[0].onclick=()=>{if(i){[files[i-1],files[i]]=[files[i],files[i-1]];renderFiles()}};bs[1].onclick=()=>{if(i<files.length-1){[files[i+1],files[i]]=[files[i],files[i+1]];renderFiles()}}}
 bs[bs.length-1].onclick=()=>{files.splice(i,1);renderFiles()};$("#files").appendChild(li)})}
function add(list){const arr=[...list].filter(f=>cur.img?/^image\//.test(f.type):/pdf$/i.test(f.type)||/\.pdf$/i.test(f.name));
 if(arr.length<list.length)msg(cur.img?T("تم تجاهل ملفات ليست صوراً.","Skipped files that are not images."):T("تم تجاهل ملفات ليست PDF.","Skipped files that are not PDFs."),"err");
 files=cur.multi?files.concat(arr):arr.slice(0,1);renderFiles()}
function msg(t,k){const m=$("#msg");m.textContent=t;m.className="msg "+(k||"")}
$("#pick").onchange=e=>add(e.target.files);
const drop=$("#drop");["dragenter","dragover"].forEach(e=>drop.addEventListener(e,ev=>{ev.preventDefault();drop.classList.add("over")}));
["dragleave","drop"].forEach(e=>drop.addEventListener(e,ev=>{ev.preventDefault();drop.classList.remove("over")}));
drop.addEventListener("drop",e=>add(e.dataTransfer.files));
document.addEventListener("dragover",e=>e.preventDefault());document.addEventListener("drop",e=>e.preventDefault());
const api=(p,body)=>fetch(p,{method:"POST",headers:{"X-Token":TOKEN,"X-Lang":L},body});
function save(blob,name){const a=document.createElement("a");a.href=URL.createObjectURL(blob);a.download=name;document.body.appendChild(a);a.click();a.remove();setTimeout(()=>URL.revokeObjectURL(a.href),5000)}
const stemOf=n=>n.replace(/\.[^.]+$/,"");
$("#go").onclick=async()=>{if(!files.length){msg(T("أضف ملفاً أولاً.","Add a file first."),"err");return}
 const fd=new FormData();fd.append("op",cur.id);
 const optv=n=>{const e=$("#opts").querySelector(`[name="${n}"]`);return e?e.value:""};
 if(cur.id==="to_text"||cur.id==="to_docx"&&optv("engine")!=="layout"){
  $("#go").disabled=true;msg(T("جاري قراءة النص…","Reading text…"));
  try{const pages=await extract(files[0],optv("range"));
   if(cur.id==="to_text"){const t=pages.map(p=>`—— ${T("صفحة","Page")} ${p.page} ——\n`+(p.lines.length?p.lines.map(l=>l.text).join("\n"):T("[لا يوجد نص]","[no text]"))).join("\n\n");
    save(new Blob(["﻿"+t],{type:"text/plain"}),stemOf(files[0].name)+".txt");msg(T("تم الحفظ: ","Saved: ")+stemOf(files[0].name)+".txt","ok");$("#go").disabled=false;return}
   const all=pages.flatMap(p=>p.lines.map(l=>l.text)).join(" "),ar=(all.match(AR_G)||[]).length,la=(all.match(/[A-Za-z]/g)||[]).length;
   if(optv("engine")==="text"||ar>=(ar+la)*0.3||!(ar+la))fd.append("lines",JSON.stringify(pages));
  }catch(e){msg(e.message||T("تعذّر قراءة الملف.","Could not read the file."),"err");$("#go").disabled=false;return}
  $("#go").disabled=false}
 $("#opts").querySelectorAll("input,select").forEach(i=>{if(i.type==="checkbox"){if(i.checked)fd.append(i.name,i.value)}else fd.append(i.name,i.value)});
 files.forEach(f=>fd.append("files",f,f.name));
 $("#go").disabled=true;msg(T("جاري المعالجة…","Working…"));
 try{const r=await api("/api/run",fd);
  if(!r.ok){const j=await r.json().catch(()=>({error:T("تعذّر التنفيذ","Could not run this tool")}));msg(j.error,"err");return}
  const name=decodeURIComponent(r.headers.get("X-Filename")||"output"),note=decodeURIComponent(r.headers.get("X-Note")||"");
  save(await r.blob(),name);msg(T("تم الحفظ: ","Saved: ")+name+(note?" — "+note:""),"ok");
 }catch(e){msg(T("انقطع الاتصال بالبرنامج. تأكد أنه ما زال يعمل.","Lost contact with the program. Check that it is still running."),"err")}finally{$("#go").disabled=false}};
$("#quit").onclick=async()=>{try{await api("/api/quit")}catch{}document.body.innerHTML=`<div class="wrap"><h1>${T("تم إيقاف البرنامج","The program has stopped")}</h1><p class="sub">${T("تقدر تغلق هذه الصفحة.","You can close this page.")}</p></div>`};
setInterval(()=>api("/api/ping").catch(()=>{}),15000);api("/api/ping");
const AR=/[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]/,AR_G=new RegExp(AR.source,"g");
const isRtl=s=>AR.test(s)&&(s.match(AR_G)||[]).length>=(s.match(/[A-Za-z]/g)||[]).length;
function parseRange(v,n){v=(v||"").trim().replace(/[٠-٩]/g,d=>"٠١٢٣٤٥٦٧٨٩".indexOf(d));if(!v)return[...Array(n).keys()];const out=[];
 for(const part of v.split(/[,،\s]+/).filter(Boolean)){const m=part.match(/^(\d+)(?:-(\d+))?$/);if(!m)throw new Error(T(`لم أفهم «${part}» في خانة الصفحات.`,`Could not read "${part}" in Pages.`));
  let a=+m[1],b=m[2]?+m[2]:a;if(a>b)[a,b]=[b,a];for(let i=a;i<=b;i++){if(i<1||i>n)throw new Error(T(`الصفحة ${i} غير موجودة؛ عدد الصفحات ${n}.`,`Page ${i} does not exist; the file has ${n} pages.`));out.push(i-1)}}return out}
function fixLine(t,rtl){if(!rtl)return t;const o=t.indexOf("("),c=t.indexOf(")");
 if(c>-1&&(o<0||c<o))t=t.replace(/[()[\]]/g,ch=>({"(":")",")":"(","[":"]","]":"["})[ch]);
 return t.replace(/\s+([ً-ْٰ])/g,"$1").replace(/(^|\s)([.,:;،؛!?])(\d[\d.,%]*)(?=\s|$)/g,"$1$3$2")}
async function pageLines(pg){const vp=pg.getViewport({scale:1}),tc=await pg.getTextContent(),items=[];
 for(const it of tc.items){if(!it.str||!it.str.trim())continue;const t=pdfjsLib.Util.transform(vp.transform,it.transform);const h=Math.hypot(t[2],t[3])||10;
  items.push({x:t[4],y:t[5],w:it.width,h,s:it.str.normalize("NFKC")})}
 items.sort((a,b)=>a.y-b.y);const lines=[];
 for(const it of items){const ln=lines[lines.length-1];if(ln&&Math.abs(it.y-ln.y)<Math.max(it.h,ln.h)*0.55){ln.items.push(it);ln.h=Math.max(ln.h,it.h)}else lines.push({y:it.y,h:it.h,items:[it]})}
 let prev=null;return lines.map(ln=>{const rtl=isRtl(ln.items.map(i=>i.s).join(" "));const arr=ln.items.sort((a,b)=>rtl?b.x-a.x:a.x-b.x);let text="";
  arr.forEach((it,k)=>{if(k){const pr=arr[k-1],gap=rtl?pr.x-(it.x+it.w):it.x-(pr.x+pr.w);if(gap>ln.h*0.12&&!/\s$/.test(text)&&!/^\s/.test(it.s))text+=" "}text+=it.s});
  const gapBefore=!!prev&&ln.y-prev.y>Math.max(prev.h,ln.h)*2.1;prev=ln;return{text:fixLine(text.replace(/\s+/g," ").trim(),rtl),rtl,size:Math.round(ln.h),gap:gapBefore}})}
async function extract(file,spec){let doc;try{doc=await pdfjsLib.getDocument({data:new Uint8Array(await file.arrayBuffer()),isEvalSupported:false}).promise}
 catch(e){throw new Error(e&&e.name==="PasswordException"?T("الملف محمي بكلمة مرور. استخدم «فك الحماية» أولاً.","This file is password-protected. Use Remove password first."):T("تعذّر فتح الملف.","Could not open the file."))}
 try{const idx=parseRange(spec,doc.numPages),out=[];for(const i of idx){msg(T(`قراءة الصفحة ${i+1}…`,`Reading page ${i+1}…`));out.push({page:i+1,lines:await pageLines(await doc.getPage(i+1))})}return out}finally{doc.destroy()}}
applyLang();
</script></body></html>"""


if __name__ == "__main__":
    if sys.version_info < (3, 9):
        sys.exit("يحتاج Python 3.9 أو أحدث. / Python 3.9 or newer is required.")
    main()
